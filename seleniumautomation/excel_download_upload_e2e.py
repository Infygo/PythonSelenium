import time
import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions

# ! upload file in Desktop > input file type and css selector

driver = webdriver.Chrome()
driver.implicitly_wait(10)
driver.get("https://rahulshettyacademy.com/upload-download-test/index.html")
driver.maximize_window()

# xpath can traverse from parent to child and child to parent
fruit_name = "Apple"
file_path = "/home/vig/Downloads/download.xlsx"
price_update = 750
price = driver.find_element(By.XPATH,
                            "//div[text()='" + fruit_name + "']/parent::div/parent::div/div[@id='cell-4-undefined']")
print("Current Price of", fruit_name, "is", price.text)

driver.find_element(By.CSS_SELECTOR, "#downloadButton").click()
time.sleep(5)


def excel_utility(file_path, fruit_name, price_update):
    # open the downloaded excel file and add some data
    open_excelbook = openpyxl.load_workbook(file_path)
    file_activesheet = open_excelbook.active
    max_colum = file_activesheet.max_column  # 5 # column numbering starts with 1
    max_row = file_activesheet.max_row  # 7 # row numbering as well starts with 1

    # Find column named 'price'
    Dict_excel = {}
    for i_row in range(1, max_row + 1):
        for j_col in range(1, max_colum + 1):
            if file_activesheet.cell(i_row, j_col).value == 'price':
                Dict_excel["price_column"] = j_col
        break
    print('Price column is :', Dict_excel["price_column"])

    # Find apple row
    for row in range(1, max_row + 1):
        for col in range(1, max_colum + 1):
            if file_activesheet.cell(row, col).value == fruit_name:
                Dict_excel["apple_row"] = row
                break
    print('Apple row is :', Dict_excel["apple_row"])
    print(Dict_excel)

    # update price for row and column
    file_activesheet.cell(row=Dict_excel['apple_row'], column=Dict_excel['price_column']).value = price_update
    open_excelbook.save(file_path)

# call function to update price for a Fruit name
excel_utility(file_path, fruit_name, price_update)

# reupload saved file
file_input = driver.find_element(By.CSS_SELECTOR, "input[type='file']")
file_input.send_keys(file_path)

# wait to see the toast message
explicit_wait = WebDriverWait(driver, 10)
toast_message_locator = (By.CSS_SELECTOR, ".Toastify__toast-body div:nth-child(2)")
